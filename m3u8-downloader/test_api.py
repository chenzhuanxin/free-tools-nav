# -*- coding: utf-8 -*-
"""m3u8tool 服务端接口回归测试

用法：
    python test_api.py [端口]
默认端口 18901。
"""
import sys
import json
import time
import urllib.parse
import urllib.request
import os

PORT = int(sys.argv[1]) if len(sys.argv) > 1 else 18901
BASE = "http://127.0.0.1:%d" % PORT

_OP = urllib.request.build_opener(urllib.request.ProxyHandler({}))

PASS = [0]
FAIL = [0]


def check(name, cond, detail=""):
    if cond:
        PASS[0] += 1
        print("  [PASS] %s" % name)
    else:
        FAIL[0] += 1
        print("  [FAIL] %s  %s" % (name, detail))


def get_json(path, timeout=40):
    # urllib 无法直接发送非 ASCII 路径，这里统一做安全编码
    safe = urllib.parse.quote(path, safe="/?&=%:,.-_~")
    with _OP.open(BASE + safe, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))


def post_json(path, obj, timeout=30):
    body = json.dumps(obj).encode("utf-8")
    req = urllib.request.Request(BASE + path, data=body,
                                 headers={"Content-Type": "application/json"})
    with _OP.open(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))


def eq(s):
    return urllib.parse.quote(s, safe="")


print("=" * 66)
print("m3u8tool 接口回归测试   目标 %s" % BASE)
print("=" * 66)

# ---------------------------------------------------------------- 1 首页
print("\n[1] 首页与静态资源")
try:
    with _OP.open(BASE + "/", timeout=15) as r:
        html = r.read().decode("utf-8")
        ct = r.headers.get("Content-Type", "")
    check("GET / 返回 200 + text/html", "text/html" in ct, ct)
    check("页面含标题", "M3U8" in html)
    check("页面含作者署名", "公歧子" in html and "gongqizi0" in html)
    check("占位符已全部替换", "__APP_NAME__" not in html and "__WECHAT__" not in html)
    check("含搜索输入框", 'id="wd"' in html)
    check("含下载队列容器", 'id="tasks"' in html)
except Exception as e:
    check("GET / ", False, str(e))
    html = ""

# ---------------------------------------------------------------- 2 配置
print("\n[2] /api/config")
try:
    c = get_json("/api/config")
    check("ok=True", c.get("ok") is True)
    check("含 sources 列表", len(c.get("sources") or []) >= 8, str(len(c.get("sources") or [])))
    check("标注了 aes_backend", bool(c.get("aes_backend")), c.get("aes_backend"))
    check("返回 out_dir", bool(c.get("out_dir")), c.get("out_dir"))
    print("       aes=%s  ffmpeg=%s" % (c.get("aes_backend"),
                                       "已探测" if c.get("ffmpeg") else "未找到"))
    check("资源站字段完整",
          all(k in s for s in c["sources"] for k in ("id", "name", "api", "search")))
except Exception as e:
    check("/api/config", False, str(e))
    c = {}

# ---------------------------------------------------------------- 3 搜索
print("\n[3] /api/search")
results = []
try:
    d = get_json("/api/search?wd=" + eq("庆余年"), timeout=60)
    results = d.get("results") or []
    check("ok=True", d.get("ok") is True)
    check("返回结果 >= 5 条", len(results) >= 5, str(len(results)))
    check("结果含标题", all(r.get("title") for r in results[:5]))
    check("结果含资源站", all(r.get("source") for r in results[:5]))
    check("结果含剧集链接",
          any(r.get("ep_count", 0) > 0 for r in results), "无任何剧集")
    withurl = [r for r in results if r.get("lines")]
    check("至少有结果带 lines", len(withurl) > 0)
    if withurl:
        eps = withurl[0]["lines"][0]["episodes"]
        check("剧集带 name/url", bool(eps) and "url" in eps[0] and "name" in eps[0])
        check("剧集 url 是 http", eps[0]["url"].startswith("http"), eps[0]["url"][:40])
    print("       共 %d 条，%d 个资源站返回" % (len(results), len(set(r["source"] for r in results))))
except Exception as e:
    check("/api/search", False, str(e))

print("\n[3b] 搜索：空关键词 / 中文乱码 / 多资源站筛选")
try:
    d = get_json("/api/search?wd=")
    check("空关键词被拒绝", d.get("ok") is False)
except Exception as e:
    check("空关键词", False, str(e))
try:
    d = get_json("/api/search?wd=" + eq("三体") + "&sources=liangzi,guangsu", timeout=40)
    srcs = set(r["source"] for r in (d.get("results") or []))
    check("指定 sources 生效", srcs.issubset({"量子资源", "光速资源"}), str(srcs))
except Exception as e:
    check("sources 筛选", False, str(e))

# ---------------------------------------------------------------- 4 解析
print("\n[4] /api/parse（m3u8 有效性检测）")
good_url = ""
for r in results:
    if r.get("lines") and r["lines"][0]["episodes"]:
        good_url = r["lines"][0]["episodes"][0]["url"]
        break
try:
    d = get_json("/api/parse?url=" + eq(good_url), timeout=40)
    check("ok=True", d.get("ok") is True, str(d)[:120])
    check("分片数 > 0", (d.get("count") or 0) > 0, str(d.get("count")))
    check("返回 kind", d.get("kind") == "media", str(d.get("kind")))
    print("       分片 %s 个，时长 %.1f 分钟，加密=%s" % (
        d.get("count"), (d.get("duration") or 0) / 60, d.get("encrypted")))
except Exception as e:
    check("/api/parse", False, str(e))

try:
    d = get_json("/api/parse?url=" + eq("https://127.0.0.1:1/nope.m3u8"), timeout=30)
    check("坏链接返回 ok=False", d.get("ok") is False)
except Exception as e:
    check("坏链接处理", False, str(e))

# ---------------------------------------------------------------- 5 任务列表
print("\n[5] /api/tasks")
try:
    d = get_json("/api/tasks")
    check("ok=True", d.get("ok") is True)
    check("tasks 是列表", isinstance(d.get("tasks"), list))
except Exception as e:
    check("/api/tasks", False, str(e))

# ---------------------------------------------------------------- 6 下载（本地小源）
print("\n[6] /api/download —— 本地 AES 加密源全流程")
try:
    from hls_fixture import make_fixture_server   # noqa
    HAVE_FIX = True
except ImportError:
    HAVE_FIX = False

if HAVE_FIX:
    srv, purl = make_fixture_server()
    t0 = time.time()
    d = get_json("/api/download?" + urllib.parse.urlencode(
        {"name": "接口测试-中文名", "url": purl, "workers": 4,
         "ffmpeg": "0", "mp4": "0"}), timeout=30)
    check("启动下载 ok=True", d.get("ok") is True, str(d)[:120])
    tid = d.get("id")
    check("中文任务名未乱码", d.get("name") == "接口测试-中文名", repr(d.get("name")))
    final = None
    for _ in range(60):
        time.sleep(0.4)
        s = get_json("/api/progress?id=" + tid)
        final = s
        if s.get("state") in ("done", "error", "canceled"):
            break
    check("下载完成", final and final.get("state") == "done",
          final and final.get("state") + " / " + str(final.get("message")))
    if final and final.get("state") == "done":
        check("分片全部成功", final.get("failed") == 0, str(final.get("failed")))
        check("产物存在", os.path.isfile(final.get("output") or ""), final.get("output"))
        sz = os.path.getsize(final["output"]) if os.path.isfile(final.get("output")) else 0
        check("产物大小 > 100KB", sz > 100 * 1024, "%d B" % sz)
        print("       用时 %.1fs，产物 %s（%d B）" % (time.time() - t0,
                                                  os.path.basename(final["output"]), sz))
    # 移除任务
    d = get_json("/api/remove?id=" + tid)
    check("移除任务 ok", d.get("ok") is True)
    check("移除后查不到", get_json("/api/progress?id=" + tid).get("ok") is False)
    srv.shutdown()
else:
    print("  [SKIP] 未找到 hls_fixture 模块")

# ---------------------------------------------------------------- 7 取消
print("\n[7] /api/cancel")
try:
    d = get_json("/api/cancel?id=" + eq("不存在的任务"))
    check("取消不存在的任务返回 ok=False", d.get("ok") is False)
except Exception as e:
    check("/api/cancel 容错", False, str(e))

# ---------------------------------------------------------------- 8 导出
print("\n[8] /api/export（xlsx / csv / json）")
rows = [{"title": r.get("title", ""), "type": r.get("type", ""), "area": r.get("area", ""),
         "year": r.get("year", ""), "ep_count": r.get("ep_count", ""),
         "source": r.get("source", ""),
         "m3u8": (r["lines"][0]["episodes"][0]["url"] if r.get("lines") and r["lines"][0]["episodes"] else ""),
         "remark": r.get("remark", ""), "updated": r.get("updated", "")}
        for r in results[:20]]
made = {}
for fmt in ("xlsx", "csv", "json"):
    try:
        d = post_json("/api/export", {"format": fmt, "wd": "接口测试导出", "rows": rows})
        made[fmt] = d.get("path")
        check("%s 导出 ok" % fmt, d.get("ok") is True, str(d)[:120])
        check("%s 文件已生成" % fmt, d.get("path") and os.path.isfile(d["path"]),
              str(d.get("path")))
        if d.get("ok") and os.path.isfile(d.get("path")):
            check("%s 大小 > 0" % fmt, os.path.getsize(d["path"]) > 0)
    except Exception as e:
        check("导出 " + fmt, False, str(e))

# xlsx 结构校验
if made.get("xlsx") and os.path.isfile(made["xlsx"]):
    import zipfile
    try:
        z = zipfile.ZipFile(made["xlsx"])
        check("xlsx 是合法 zip", z.testzip() is None)
        names = z.namelist()
        check("xlsx 含 sheet1.xml", "xl/worksheets/sheet1.xml" in names, str(names))
        check("xlsx 含 styles.xml", "xl/styles.xml" in names)
        sheet = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        check("xlsx 含冻结窗格", 'ySplit="1"' in sheet)
        check("xlsx 含自动筛选", "<autoFilter" in sheet)
        check("xlsx 表头含序号", "序号" in sheet and "m3u8链接" in sheet)
        try:
            import openpyxl
            wb = openpyxl.load_workbook(made["xlsx"])
            ws = wb.active
            check("openpyxl 可读取", ws.max_row == len(rows) + 1,
                  "行数 %s 期望 %s" % (ws.max_row, len(rows) + 1))
            hdr = [cell.value for cell in ws[1]]
            check("表头 10 列正确", hdr[:3] == ["序号", "标题", "类型"], str(hdr))
        except ImportError:
            print("       [SKIP] 无 openpyxl，跳过交叉校验")
    except Exception as e:
        check("xlsx 结构", False, str(e))

# ---------------------------------------------------------------- 9 异常输入
print("\n[9] 异常输入容错")
for p in ("/api/progress", "/api/cancel", "/api/remove", "/api/open"):
    try:
        d = get_json(p)
        check("%s 无参数不崩溃" % p, isinstance(d, dict))
    except Exception as e:
        check("%s 无参数" % p, False, str(e))
try:
    with _OP.open(BASE + "/" + urllib.parse.quote("不存在的路径"), timeout=10) as r:
        check("未知路径返回 404", r.status == 404, str(r.status))
except urllib.error.HTTPError as e:
    check("未知路径返回 404", e.code == 404, str(e.code))
except Exception as e:
    check("未知路径", False, str(e))

try:
    d = get_json("/api/download?url=" + eq("https://127.0.0.1:1/x.m3u8") + "&name=坏源")
    tid = d.get("id")
    final = None
    for _ in range(40):
        time.sleep(0.4)
        final = get_json("/api/progress?id=" + tid)
        if final.get("state") in ("done", "error", "canceled"):
            break
    check("坏源最终 state=error", final and final.get("state") == "error",
          final and final.get("state"))
    check("坏源有错误信息", bool(final and final.get("error")))
except Exception as e:
    check("坏源处理", False, str(e))

# ---------------------------------------------------------------- 汇总
print("\n" + "=" * 66)
print("结果：%d 通过，%d 失败" % (PASS[0], FAIL[0]))
print("=" * 66)
sys.exit(1 if FAIL[0] else 0)
